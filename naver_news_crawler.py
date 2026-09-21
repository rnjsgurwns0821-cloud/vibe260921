"""
네이버 통합검색 결과에서 뉴스 기사 링크를 수집하고,
각 기사(n.news.naver.com)의 제목/언론사/기자/작성일/본문을 크롤링하는 스크립트.

필요 패키지: requests, beautifulsoup4, openpyxl (모두 requirements.txt에 포함되어 있음)
    pip install -r requirements.txt
"""

import csv
import re
import time
from urllib.parse import parse_qs, urlparse

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
    )
}

# 크롤링할 네이버 통합검색 URL (질의어: 반도체)
SEARCH_URL = (
    "https://search.naver.com/search.naver"
    "?where=nexearch&sm=top_hty&fbm=0&ie=utf8&query=%EB%B0%98%EB%8F%84%EC%B2%B4"
)

ARTICLE_LINK_PATTERN = re.compile(r"https://n\.news\.naver\.com/mnews/article/\d+/\d+")


def get_article_links(search_url):
    """검색 결과 페이지에서 네이버 뉴스 기사 링크를 중복 없이 추출한다."""
    res = requests.get(search_url, headers=HEADERS, timeout=10)
    res.raise_for_status()
    res.encoding = res.apparent_encoding

    soup = BeautifulSoup(res.text, "html.parser")

    links = []
    seen = set()
    for a_tag in soup.find_all("a", href=True):
        href = a_tag["href"]
        if ARTICLE_LINK_PATTERN.match(href):
            clean_url = href.split("?")[0]
            if clean_url not in seen:
                seen.add(clean_url)
                links.append(href)

    return links


def _extract_real_image_url(proxied_src):
    """search.pstatic.net 프록시 썸네일 URL에서 원본 이미지 URL을 추출한다."""
    if not proxied_src:
        return ""
    query = parse_qs(urlparse(proxied_src).query)
    return query.get("src", [proxied_src])[0]


def parse_search_result_page(html):
    """네이버 통합검색 뉴스 결과 목록 HTML에서 기사 미리보기 정보를 추출한다.

    개별 기사 페이지에 접속하지 않고 검색결과 리스트에 이미 표시된 정보만 사용하므로
    제목/언론사/링크/상대시간/요약/썸네일은 얻을 수 있지만 기자명과 본문 전체는 얻을 수 없다.
    각 기사 그룹의 대표 기사는 level="main", 하위 관련기사는 level="related"로 구분된다.
    """
    soup = BeautifulSoup(html, "html.parser")

    results = []
    for profile in soup.select('div[data-sds-comp="Profile"]'):
        group = profile.parent
        if group is None:
            continue

        is_main = "title-color-g10" in profile.get("class", [])
        area_prefix = "nws_all.h." if is_main else "nws_all.i."

        title_tag = group.select_one(f'a[data-nlog-area="{area_prefix}tit"]')
        if not title_tag or not title_tag.get("href"):
            continue

        press_tag = profile.select_one(".sds-comps-profile-info-title-text")
        subtexts = profile.select(".sds-comps-profile-info-subtext")
        nav_tag = profile.select_one(f'a[data-nlog-area="{area_prefix}nav"]')

        snippet = ""
        thumbnail = ""
        if is_main:
            body_tag = group.select_one(f'a[data-nlog-area="{area_prefix}body"]')
            snippet = body_tag.get_text(strip=True) if body_tag else ""

            img_tag = group.select_one(f'a[data-nlog-area="{area_prefix}img"] img')
            if img_tag and img_tag.get("src"):
                thumbnail = _extract_real_image_url(img_tag["src"])

        results.append(
            {
                "level": "main" if is_main else "related",
                "title": title_tag.get_text(strip=True),
                "press": press_tag.get_text(strip=True) if press_tag else "",
                "url": title_tag["href"],
                "naver_url": nav_tag["href"] if nav_tag and nav_tag.get("href") else "",
                "published": subtexts[0].get_text(strip=True) if subtexts else "",
                "snippet": snippet,
                "thumbnail": thumbnail,
            }
        )

    return results


def get_preview_articles(search_url):
    """검색결과 페이지에 한 번만 접속하여 기사 미리보기 목록을 가져온다 (개별 기사 크롤링 없음)."""
    res = requests.get(search_url, headers=HEADERS, timeout=10)
    res.raise_for_status()
    res.encoding = res.apparent_encoding
    return parse_search_result_page(res.text)


def parse_article(article_url):
    """네이버 뉴스 기사 페이지에서 제목/언론사/기자/작성일/본문을 추출한다."""
    res = requests.get(article_url, headers=HEADERS, timeout=10)
    res.raise_for_status()
    res.encoding = res.apparent_encoding

    soup = BeautifulSoup(res.text, "html.parser")

    title_tag = soup.select_one("#title_area span")
    title = title_tag.get_text(strip=True) if title_tag else ""

    press_tag = soup.select_one(".media_end_head_top_logo img")
    press = press_tag.get("alt", "").strip() if press_tag else ""

    reporter_tag = soup.select_one(".media_end_head_journalist_name")
    reporter = reporter_tag.get_text(strip=True) if reporter_tag else ""

    date_tag = soup.select_one(".media_end_head_info_datestamp_time")
    published_at = date_tag.get("data-date-time", "") if date_tag else ""

    body_tag = soup.select_one("#dic_area")
    if body_tag:
        for unwanted in body_tag.select("script, style, .end_photo_org, .vod_player_wrap"):
            unwanted.decompose()
        content = body_tag.get_text("\n", strip=True)
    else:
        content = ""

    return {
        "url": article_url,
        "title": title,
        "press": press,
        "reporter": reporter,
        "published_at": published_at,
        "content": content,
    }


def crawl(search_url, delay=0.5):
    """검색 결과의 기사 링크를 모두 크롤링하여 리스트로 반환한다."""
    links = get_article_links(search_url)
    print(f"[INFO] 수집된 기사 링크 수: {len(links)}")

    articles = []
    for i, link in enumerate(links, start=1):
        try:
            article = parse_article(link)
            articles.append(article)
            print(f"[{i}/{len(links)}] {article['title']} ({article['press']})")
        except requests.RequestException as e:
            print(f"[WARN] 크롤링 실패: {link} ({e})")
        time.sleep(delay)  # 서버 부하 방지를 위한 딜레이

    return articles


def save_to_csv(articles, filename="naver_news_result.csv"):
    with open(filename, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(
            f, fieldnames=["title", "press", "reporter", "published_at", "url", "content"]
        )
        writer.writeheader()
        writer.writerows(articles)
    print(f"[INFO] 결과 저장 완료: {filename}")


def save_to_excel(articles, filename="naver_news_result.xlsx"):
    """크롤링 결과를 엑셀(.xlsx) 파일로 저장한다."""
    headers = ["제목", "언론사", "기자", "작성일", "URL", "본문"]
    field_order = ["title", "press", "reporter", "published_at", "url", "content"]
    column_widths = [45, 12, 12, 20, 45, 80]

    wb = Workbook()
    ws = wb.active
    ws.title = "네이버뉴스"

    ws.append(headers)
    for col, width in enumerate(column_widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"

    for article in articles:
        row = [article.get(field, "") for field in field_order]
        ws.append(row)
        row_idx = ws.max_row
        ws.cell(row=row_idx, column=6).alignment = Alignment(wrap_text=True, vertical="top")
        for col in (1, 2, 3, 4, 5):
            ws.cell(row=row_idx, column=col).alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(filename)
    print(f"[INFO] 결과 저장 완료: {filename}")


if __name__ == "__main__":
    result = crawl(SEARCH_URL)
    save_to_csv(result)
